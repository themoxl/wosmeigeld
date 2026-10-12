"""DB-Orchestrierung und Normalisierung fuer den Excel-/CSV-Kostenplanimport
(Plan Abschnitt 10). Sicherheitspruefungen leben in
`project_cost_import_security.py`, die Kostenabschnitts-/-positions-Domain in
`project_service.py`; dieser Dienst verbindet beides: Datei -> normalisierte
Zeilen -> Preview mit Fingerprint -> Apply mit Idempotenz.

Zwei bewusste Architekturentscheidungen (Plan 10.2):

1. Sowohl Preview als auch Apply parsen die HOCHGELADENE DATEI selbst neu,
   statt dem Client die einmal geparsten Zeilen zurueckzugeben und beim Apply
   ungeprueft zu uebernehmen. Apply prueft den `preview_fingerprint` gegen das
   frische Parse-Ergebnis - stimmt er nicht mehr (Datei geaendert, Mapping
   geaendert, oder eine bereits zugeordnete Position wurde inzwischen von
   jemand anderem bearbeitet), scheitert Apply mit 409 statt eine veraltete
   Preview blind anzuwenden.
2. "Niemals still anhand aehnlicher Namen ueberschreiben" (Plan 10.2, Schritt
   7) heisst konkret: die Preview SCHLAEGT eine Zuordnung zu einer
   bestehenden Position per Namensgleichheit nur VOR (`matched_item_id`).
   Welche Zeile tatsaechlich eine bestehende Position aktualisiert, welche
   eine neue anlegt und welche uebersprungen wird, entscheidet ausschliesslich
   die vom Aufrufer beim Apply mitgelieferte, explizite Zeilenliste
   (`CostImportRowResolution`). Der Server trifft diese Entscheidung nie
   selbst."""

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Final, cast
from uuid import UUID

import openpyxl
from fastapi import HTTPException, status
from sqlalchemy import text
from sqlalchemy.engine import CursorResult
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.project_import import (
    CostImportApplyResult,
    CostImportMappingProfile,
    CostImportMappingProfileInput,
    CostImportPreviewResponse,
    CostImportPreviewRow,
    CostImportPreviewSums,
    CostImportRowResolution,
    ImportFileFormat,
)
from app.schemas.projects import CostItemCreate, CostItemUpdate, CostSectionCreate
from app.services import project_service
from app.services.audit_service import record_audit_event
from app.services.project_cost_import_security import (
    CostImportSecurityError,
    check_file_size,
    check_row_count,
    check_sheet_count,
    check_xlsx_structure,
)

COST_IMPORT_PROFILE_NOT_FOUND: Final = "cost_import_profile_not_found"
PREVIEW_FINGERPRINT_MISMATCH: Final = "preview_fingerprint_mismatch"
UNKNOWN_ROW_RESOLUTION: Final = "unknown_row_resolution"

_STATUS_VALUES: Final = frozenset(
    {"estimated", "quoted", "committed", "partially_paid", "paid", "cancelled"}
)

# Generische deutsche/englische Spaltenaliase (Plan 10.1). Nur zur
# Vorbelegung eines Mappingvorschlags im Web - massgeblich ist immer das vom
# Aufrufer uebergebene `column_mapping`, nicht diese Liste.
COLUMN_ALIASES: Final[dict[str, tuple[str, ...]]] = {
    "section_name": ("abschnitt", "section", "bereich"),
    "item_name": ("position", "item", "name", "bezeichnung"),
    "status": ("status",),
    "quantity": ("menge", "quantity", "anzahl"),
    "unit_label": ("einheit", "unit"),
    "unit_price_minor": ("einzelpreis", "unit_price", "preis_pro_einheit", "preis"),
    "estimate_manual_minor": ("schätzung", "schaetzung", "estimate", "kalkulation"),
    "quoted_amount_minor": ("angebot", "quote", "angebotspreis"),
    "committed_amount_minor": ("beauftragt", "committed", "auftragssumme"),
    "expected_start_on": ("start", "beginn", "startdatum"),
    "expected_end_on": ("ende", "end", "enddatum"),
    "vendor_label": ("anbieter", "vendor", "lieferant", "firma"),
    "note": ("notiz", "note", "bemerkung", "anmerkung"),
}

_DATE_FORMATS: Final = ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y")
_AMOUNT_CLEAN_RE: Final = re.compile(r"[^\d,.\-]")


@dataclass
class _NormalizedRow:
    row_number: int
    section_name: str | None = None
    item_name: str | None = None
    status: str | None = None
    quantity: Decimal | None = None
    unit_label: str | None = None
    unit_price_minor: int | None = None
    estimate_manual_minor: int | None = None
    quoted_amount_minor: int | None = None
    committed_amount_minor: int | None = None
    vendor_label: str | None = None
    expected_start_on: date | None = None
    expected_end_on: date | None = None
    note: str | None = None
    errors: list[str] = field(default_factory=list)

    @property
    def calculation_mode(self) -> str:
        if self.quantity is not None and self.unit_price_minor is not None:
            return "quantity"
        return "fixed"


def _parse_amount_minor(raw: str) -> int:
    """Deutsches (1.234,56) und englisches (1,234.56 / 1234.56) Format sowie
    reines Komma- oder Punkt-Dezimaltrennzeichen (Plan 15.5: "Dezimalkomma/
    -punkt")."""
    cleaned = _AMOUNT_CLEAN_RE.sub("", raw.strip())
    if not cleaned or cleaned == "-":
        raise ValueError("leerer Betrag")
    negative = cleaned.startswith("-")
    cleaned = cleaned.lstrip("-")
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    value = Decimal(cleaned)
    minor = int((value * 100).to_integral_value(rounding=ROUND_HALF_UP))
    return -minor if negative else minor


def _parse_date(raw: Any) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text_val = str(raw).strip()
    if not text_val:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text_val, fmt).date()
        except ValueError:
            continue
    return None


def _clean_str(value: Any) -> str | None:
    if value is None:
        return None
    text_val = str(value).strip()
    return text_val or None


def _detect_format(filename: str) -> ImportFileFormat:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return "csv_generic"
    if lower.endswith(".xlsx"):
        return "xlsx_generic"
    raise CostImportSecurityError("Nur .xlsx und .csv werden unterstuetzt.")


def _read_csv_rows(file_bytes: bytes) -> tuple[list[str], list[list[str]]]:
    text_content = file_bytes.decode("utf-8-sig")
    sample = text_content[:4096]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=";,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.reader(io.StringIO(text_content), dialect)
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise CostImportSecurityError("Datei enthaelt keine Daten.")
    headers = [h.strip() for h in rows[0]]
    check_row_count(len(rows) - 1)
    return headers, rows[1:]


def _read_xlsx_rows(
    file_bytes: bytes, sheet_name: str | None, header_row: int
) -> tuple[list[str], list[list[Any]], list[str]]:
    # check_xlsx_structure() prueft die ZIP-Huelle strukturell vorab, aber ein
    # Archiv kann diese Pruefung bestehen und trotzdem kein gueltiges,
    # vollstaendiges XLSX sein (z. B. wenn Metadaten in der ZIP-Zentral-
    # verzeichnisstruktur zum tatsaechlichen komprimierten Inhalt nicht
    # passen). Ohne dieses except waere das ein unbehandelter 500 statt eines
    # klaren 400.
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except CostImportSecurityError:
        raise
    except Exception as exc:
        raise CostImportSecurityError(
            "Datei ist kein gueltiges XLSX (moeglicherweise beschaedigt)."
        ) from exc
    try:
        sheet_names = list(wb.sheetnames)
        check_sheet_count(len(sheet_names))
        sheet = wb[sheet_name] if sheet_name else wb.active
        if sheet is None:
            raise CostImportSecurityError("Kein aktives Blatt in der Datei.")
        all_rows = list(sheet.iter_rows(values_only=True))
        if header_row < 1 or header_row > len(all_rows):
            raise CostImportSecurityError(f"Headerzeile {header_row} existiert nicht.")
        headers = [str(c).strip() if c is not None else "" for c in all_rows[header_row - 1]]
        data_rows = [
            list(r) for r in all_rows[header_row:] if any(c is not None for c in r)
        ]
        check_row_count(len(data_rows))
        return headers, data_rows, sheet_names
    finally:
        wb.close()


def _parse_file(
    file_bytes: bytes, filename: str, sheet_name: str | None, header_row: int
) -> tuple[ImportFileFormat, list[str], list[list[Any]], list[str]]:
    check_file_size(file_bytes)
    detected_format = _detect_format(filename)
    if detected_format == "xlsx_generic":
        check_xlsx_structure(file_bytes, filename)
        headers, data_rows, sheet_names = _read_xlsx_rows(file_bytes, sheet_name, header_row)
    else:
        headers, data_rows = _read_csv_rows(file_bytes)
        sheet_names = []
    return detected_format, headers, data_rows, sheet_names


def _row_to_dict(headers: list[str], row: list[Any]) -> dict[str, Any]:
    return {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}


def _normalize_row(
    row_number: int, raw: dict[str, Any], mapping: dict[str, str]
) -> _NormalizedRow:
    normalized = _NormalizedRow(row_number=row_number)

    def get(field_key: str) -> Any:
        col = mapping.get(field_key)
        if not col:
            return None
        return raw.get(col)

    normalized.section_name = _clean_str(get("section_name"))
    normalized.item_name = _clean_str(get("item_name"))

    status_raw = _clean_str(get("status"))
    if status_raw:
        status_lower = status_raw.lower()
        if status_lower in _STATUS_VALUES:
            normalized.status = status_lower
        else:
            normalized.errors.append(f"unbekannter Status ignoriert: {status_raw!r}")

    quantity_raw = _clean_str(get("quantity"))
    if quantity_raw:
        try:
            normalized.quantity = Decimal(quantity_raw.replace(",", "."))
        except InvalidOperation:
            normalized.errors.append(f"ungueltige Menge: {quantity_raw!r}")

    normalized.unit_label = _clean_str(get("unit_label"))
    normalized.vendor_label = _clean_str(get("vendor_label"))
    normalized.note = _clean_str(get("note"))

    if (raw_val := _clean_str(get("unit_price_minor"))) is not None:
        try:
            normalized.unit_price_minor = _parse_amount_minor(raw_val)
        except (InvalidOperation, ValueError):
            normalized.errors.append(f"ungueltiger Einzelpreis: {raw_val!r}")
    if (raw_val := _clean_str(get("estimate_manual_minor"))) is not None:
        try:
            normalized.estimate_manual_minor = _parse_amount_minor(raw_val)
        except (InvalidOperation, ValueError):
            normalized.errors.append(f"ungueltige Schaetzung: {raw_val!r}")
    if (raw_val := _clean_str(get("quoted_amount_minor"))) is not None:
        try:
            normalized.quoted_amount_minor = _parse_amount_minor(raw_val)
        except (InvalidOperation, ValueError):
            normalized.errors.append(f"ungueltiges Angebot: {raw_val!r}")
    if (raw_val := _clean_str(get("committed_amount_minor"))) is not None:
        try:
            normalized.committed_amount_minor = _parse_amount_minor(raw_val)
        except (InvalidOperation, ValueError):
            normalized.errors.append(f"ungueltiger Beauftragt-Betrag: {raw_val!r}")

    for field_key, label in (("expected_start_on", "Start"), ("expected_end_on", "Ende")):
        raw_val_date = get(field_key)
        if raw_val_date is None or (isinstance(raw_val_date, str) and not raw_val_date.strip()):
            continue
        parsed = _parse_date(raw_val_date)
        if parsed is None:
            normalized.errors.append(f"ungueltiges Datum in {label}: {raw_val_date!r}")
        elif field_key == "expected_start_on":
            normalized.expected_start_on = parsed
        else:
            normalized.expected_end_on = parsed

    if not normalized.item_name:
        normalized.errors.append("Position ohne Namen")

    return normalized


def _compute_fingerprint(
    column_mapping: dict[str, str],
    sheet_name: str | None,
    header_row: int,
    rows: list[_NormalizedRow],
    section_versions: dict[str, int],
    item_versions: dict[str, int],
) -> str:
    payload = {
        "column_mapping": dict(sorted(column_mapping.items())),
        "sheet_name": sheet_name,
        "header_row": header_row,
        "rows": [
            {
                "row_number": r.row_number,
                "section_name": r.section_name,
                "item_name": r.item_name,
                "status": r.status,
                "quantity": str(r.quantity) if r.quantity is not None else None,
                "unit_label": r.unit_label,
                "unit_price_minor": r.unit_price_minor,
                "estimate_manual_minor": r.estimate_manual_minor,
                "quoted_amount_minor": r.quoted_amount_minor,
                "committed_amount_minor": r.committed_amount_minor,
                "vendor_label": r.vendor_label,
                "expected_start_on": (
                    r.expected_start_on.isoformat() if r.expected_start_on else None
                ),
                "expected_end_on": r.expected_end_on.isoformat() if r.expected_end_on else None,
                "note": r.note,
                "errors": r.errors,
            }
            for r in rows
        ],
        "section_versions": dict(sorted(section_versions.items())),
        "item_versions": dict(sorted(item_versions.items())),
    }
    encoded = json.dumps(payload, sort_keys=True, default=str).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


async def _build_preview_with_versions(
    session: AsyncSession,
    household_id: UUID,
    project_id: UUID,
    scenario_id: UUID,
    *,
    file_bytes: bytes,
    filename: str,
    column_mapping: dict[str, str],
    sheet_name: str | None,
    header_row: int,
) -> tuple[CostImportPreviewResponse, dict[str, int]]:
    """Wie `build_preview()`, gibt zusaetzlich die beim Matching gelesenen
    Positionsversionen zurueck. `apply_import()` nutzt sie als
    `expected_version`, damit ein zwischen diesem Read und dem eigentlichen
    Schreiben spaeter im selben Request geaenderter Datensatz einen 409
    ausloest statt (unter READ COMMITTED) unbemerkt ueberschrieben zu werden
    -- der oeffentliche Preview-Fingerprint allein deckt nur die Zeit
    zwischen Preview- und Apply-AUFRUF ab, nicht die Zeit innerhalb eines
    einzelnen Apply-Laufs mit vielen Zeilen."""
    await project_service.require_scenario(session, household_id, project_id, scenario_id)

    try:
        detected_format, headers, data_rows, sheet_names = _parse_file(
            file_bytes, filename, sheet_name, header_row
        )
    except CostImportSecurityError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc

    normalized_rows = [
        _normalize_row(index + 1, _row_to_dict(headers, row), column_mapping)
        for index, row in enumerate(data_rows)
    ]

    existing_sections = await project_service.list_cost_sections(
        session, household_id, scenario_id
    )
    existing_items = await project_service.list_cost_items(session, household_id, scenario_id)
    sections_by_name = {s.name.strip().lower(): s for s in existing_sections}
    items_by_name = {i.name.strip().lower(): i for i in existing_items}

    preview_rows: list[CostImportPreviewRow] = []
    section_versions: dict[str, int] = {}
    item_versions: dict[str, int] = {}
    estimate_total = quoted_total = committed_total = 0
    new_items_count = matched_items_count = rows_with_errors_count = 0

    for row in normalized_rows:
        matched_section = (
            sections_by_name.get(row.section_name.strip().lower()) if row.section_name else None
        )
        matched_item = items_by_name.get(row.item_name.strip().lower()) if row.item_name else None

        if matched_section is not None:
            section_versions[str(matched_section.id)] = matched_section.version
        if matched_item is not None:
            item_versions[str(matched_item.id)] = matched_item.version
            matched_items_count += 1
        elif row.item_name:
            new_items_count += 1

        if row.errors:
            rows_with_errors_count += 1
        else:
            estimate_total += row.estimate_manual_minor or 0
            quoted_total += row.quoted_amount_minor or 0
            committed_total += row.committed_amount_minor or 0

        preview_rows.append(
            CostImportPreviewRow(
                row_number=row.row_number,
                section_name=row.section_name,
                item_name=row.item_name,
                status=row.status,
                calculation_mode=row.calculation_mode,  # type: ignore[arg-type]
                quantity=row.quantity,
                unit_label=row.unit_label,
                unit_price_minor=row.unit_price_minor,
                estimate_manual_minor=row.estimate_manual_minor,
                quoted_amount_minor=row.quoted_amount_minor,
                committed_amount_minor=row.committed_amount_minor,
                vendor_label=row.vendor_label,
                expected_start_on=row.expected_start_on,
                expected_end_on=row.expected_end_on,
                note=row.note,
                matched_section_id=matched_section.id if matched_section else None,
                matched_item_id=matched_item.id if matched_item else None,
                errors=row.errors,
            )
        )

    fingerprint = _compute_fingerprint(
        column_mapping, sheet_name, header_row, normalized_rows, section_versions, item_versions
    )

    response = CostImportPreviewResponse(
        preview_fingerprint=fingerprint,
        detected_format=detected_format,
        sheet_names=sheet_names,
        selected_sheet=sheet_name,
        rows=preview_rows,
        sums=CostImportPreviewSums(
            row_count=len(preview_rows),
            rows_with_errors_count=rows_with_errors_count,
            new_items_count=new_items_count,
            matched_items_count=matched_items_count,
            estimate_total_minor=estimate_total,
            quoted_total_minor=quoted_total,
            committed_total_minor=committed_total,
        ),
    )
    return response, item_versions


async def build_preview(
    session: AsyncSession,
    household_id: UUID,
    project_id: UUID,
    scenario_id: UUID,
    *,
    file_bytes: bytes,
    filename: str,
    column_mapping: dict[str, str],
    sheet_name: str | None,
    header_row: int,
) -> CostImportPreviewResponse:
    response, _item_versions = await _build_preview_with_versions(
        session,
        household_id,
        project_id,
        scenario_id,
        file_bytes=file_bytes,
        filename=filename,
        column_mapping=column_mapping,
        sheet_name=sheet_name,
        header_row=header_row,
    )
    return response


def _row_to_create_payload(row: CostImportPreviewRow, section_id: UUID) -> CostItemCreate:
    assert row.item_name is not None
    kwargs: dict[str, Any] = {
        "section_id": section_id,
        "name": row.item_name,
        "vendor_label": row.vendor_label,
        "expected_start_on": row.expected_start_on,
        "expected_end_on": row.expected_end_on,
        "note": row.note,
        "quoted_amount_minor": row.quoted_amount_minor,
        "committed_amount_minor": row.committed_amount_minor,
    }
    if row.status:
        kwargs["status"] = row.status
    if row.calculation_mode == "quantity":
        kwargs["calculation_mode"] = "quantity"
        kwargs["quantity"] = row.quantity
        kwargs["unit_label"] = row.unit_label
        kwargs["unit_price_minor"] = row.unit_price_minor
    else:
        kwargs["calculation_mode"] = "fixed"
        # CostItemCreate verlangt bei calculation_mode=fixed eine Zahl (auch 0)
        # - ein Platzhalter ohne jede Schaetzung darf nicht am Pydantic-
        # Validator scheitern, sondern legt eine Position mit 0 EUR an, die
        # der Nutzer danach befuellt.
        kwargs["estimate_manual_minor"] = row.estimate_manual_minor or 0
    return CostItemCreate(**kwargs)


def _row_to_update_payload(row: CostImportPreviewRow, section_id: UUID) -> CostItemUpdate:
    """Setzt nur Felder, die die Importzeile tatsaechlich mitbringt - eine
    Spalte, die in der Datei leer ist, loescht kein bestehendes Feld der
    Zielposition (kein stilles Ueberschreiben durch unvollstaendige
    Importdaten, Plan 10.2)."""
    kwargs: dict[str, Any] = {"section_id": section_id}
    if row.item_name:
        kwargs["name"] = row.item_name
    if row.status:
        kwargs["status"] = row.status
    if row.vendor_label is not None:
        kwargs["vendor_label"] = row.vendor_label
    if row.expected_start_on is not None:
        kwargs["expected_start_on"] = row.expected_start_on
    if row.expected_end_on is not None:
        kwargs["expected_end_on"] = row.expected_end_on
    if row.note is not None:
        kwargs["note"] = row.note
    if row.quoted_amount_minor is not None:
        kwargs["quoted_amount_minor"] = row.quoted_amount_minor
    if row.committed_amount_minor is not None:
        kwargs["committed_amount_minor"] = row.committed_amount_minor
    if (
        row.calculation_mode == "quantity"
        and row.quantity is not None
        and row.unit_price_minor is not None
    ):
        kwargs["calculation_mode"] = "quantity"
        kwargs["quantity"] = row.quantity
        kwargs["unit_price_minor"] = row.unit_price_minor
        if row.unit_label is not None:
            kwargs["unit_label"] = row.unit_label
    elif row.estimate_manual_minor is not None:
        kwargs["calculation_mode"] = "fixed"
        kwargs["estimate_manual_minor"] = row.estimate_manual_minor
        if row.unit_label is not None:
            kwargs["unit_label"] = row.unit_label
    return CostItemUpdate(**kwargs)


async def apply_import(
    session: AsyncSession,
    household_id: UUID,
    project_id: UUID,
    scenario_id: UUID,
    user_id: UUID,
    *,
    file_bytes: bytes,
    filename: str,
    column_mapping: dict[str, str],
    sheet_name: str | None,
    header_row: int,
    resolutions: list[CostImportRowResolution],
    request_id: str,
    preview_fingerprint: str,
) -> CostImportApplyResult:
    await project_service.require_scenario(session, household_id, project_id, scenario_id)

    # scenario_id ist Teil des Idempotenzschluessels: eine request_id ist nur
    # innerhalb DESSELBEN Szenarios eindeutig - sonst koennte ein Client, der
    # denselben request_id versehentlich fuer zwei Szenarien desselben
    # Projekts verwendet, beim zweiten Aufruf faelschlich das Ergebnis des
    # ersten zurueckbekommen, ohne dass der zweite Import je gelaufen ist.
    existing = (
        await session.execute(
            text(
                "select result_json from project_cost_import_applies "
                "where household_id = :household_id and project_id = :project_id "
                "and scenario_id = :scenario_id and request_id = :request_id"
            ),
            {
                "household_id": str(household_id),
                "project_id": str(project_id),
                "scenario_id": str(scenario_id),
                "request_id": request_id,
            },
        )
    ).first()
    if existing is not None:
        result = CostImportApplyResult(**existing.result_json)
        return result.model_copy(update={"idempotent_replay": True})

    preview, item_versions = await _build_preview_with_versions(
        session,
        household_id,
        project_id,
        scenario_id,
        file_bytes=file_bytes,
        filename=filename,
        column_mapping=column_mapping,
        sheet_name=sheet_name,
        header_row=header_row,
    )
    if preview.preview_fingerprint != preview_fingerprint:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail=PREVIEW_FINGERPRINT_MISMATCH
        )

    rows_by_number = {row.row_number: row for row in preview.rows}
    resolutions_by_number = {r.row_number: r for r in resolutions}
    if set(rows_by_number) != set(resolutions_by_number):
        raise HTTPException(status_code=422, detail=UNKNOWN_ROW_RESOLUTION)

    created_ids: list[UUID] = []
    updated_ids: list[UUID] = []
    skipped_numbers: list[int] = []
    # Zeilen mit gleichem, noch unbekanntem Abschnittsnamen teilen sich einen
    # neu angelegten Abschnitt, statt pro Zeile einen Duplikat-Abschnitt zu
    # erzeugen (die Preview kennt matched_section_id nur fuer BEREITS
    # bestehende Abschnitte, nicht fuer welche, die dieser Apply-Lauf selbst
    # gerade erst anlegt).
    created_sections_by_name: dict[str, UUID] = {}

    for row_number, resolution in resolutions_by_number.items():
        row = rows_by_number[row_number]
        if resolution.action == "skip" or row.errors:
            skipped_numbers.append(row_number)
            continue

        section_id = row.matched_section_id
        if section_id is None:
            if not row.section_name:
                skipped_numbers.append(row_number)
                continue
            section_key = row.section_name.strip().lower()
            section_id = created_sections_by_name.get(section_key)
            if section_id is None:
                new_section = await project_service.create_cost_section(
                    session, household_id, scenario_id, CostSectionCreate(name=row.section_name)
                )
                section_id = new_section.id
                created_sections_by_name[section_key] = section_id

        if resolution.action == "update":
            if resolution.target_item_id is None:
                raise HTTPException(status_code=422, detail=UNKNOWN_ROW_RESOLUTION)
            # target_item_id kommt vom Aufrufer und darf NIE blind vertraut
            # werden: require_cost_item prueft schon die Haushaltsgrenze,
            # zusaetzlich muss die Position zum gerade importierten Szenario
            # gehoeren - sonst koennte eine Zeile aus Projekt/Szenario A eine
            # Position in Projekt/Szenario B desselben Haushalts ueberschreiben.
            target_item = await project_service.require_cost_item(
                session, household_id, resolution.target_item_id
            )
            if target_item.scenario_id != scenario_id:
                raise HTTPException(status_code=422, detail=UNKNOWN_ROW_RESOLUTION)
            update_payload = _row_to_update_payload(row, section_id)
            # expected_version stammt aus DEMSELBEN Preview-Read wie der
            # Fingerprint (nicht aus einem frischen Read direkt vor dem
            # Update) - so faellt eine Aenderung auf, die zwischen dem
            # Preview-Read dieses Apply-Laufs und dem Erreichen dieser Zeile
            # in der Schleife committet wurde, statt sie stillschweigend zu
            # ueberschreiben.
            await project_service.update_cost_item(
                session,
                household_id,
                resolution.target_item_id,
                update_payload,
                # Bevorzugt die beim Preview-Read erfasste Version (deckt den
                # gesamten Zeitraum seit der Vorschau ab); faellt fuer ein vom
                # Aufrufer manuell statt per Namensvorschlag gewaehltes Ziel
                # (dort kennt die Preview keine Version) auf die soeben
                # gelesene aktuelle Version zurueck, statt die Pruefung ganz
                # auszulassen.
                item_versions.get(str(resolution.target_item_id), target_item.version),
            )
            updated_ids.append(resolution.target_item_id)
        else:
            create_payload = _row_to_create_payload(row, section_id)
            created = await project_service.create_cost_item(
                session, household_id, scenario_id, user_id, create_payload
            )
            created_ids.append(created.id)

    result = CostImportApplyResult(
        request_id=request_id,
        idempotent_replay=False,
        created_item_ids=created_ids,
        updated_item_ids=updated_ids,
        skipped_row_numbers=skipped_numbers,
    )

    await session.execute(
        text(
            "insert into project_cost_import_applies "
            "(household_id, project_id, scenario_id, request_id, preview_fingerprint, "
            "result_json) "
            "values (:household_id, :project_id, :scenario_id, :request_id, "
            ":preview_fingerprint, cast(:result_json as jsonb))"
        ),
        {
            "household_id": str(household_id),
            "project_id": str(project_id),
            "scenario_id": str(scenario_id),
            "request_id": request_id,
            "preview_fingerprint": preview_fingerprint,
            "result_json": result.model_dump_json(),
        },
    )
    await record_audit_event(
        session,
        household_id,
        user_id,
        "project_cost_import",
        project_id,
        "create",
        after=result.model_dump(mode="json"),
    )
    return result


async def list_mapping_profiles(
    session: AsyncSession, household_id: UUID
) -> list[CostImportMappingProfile]:
    rows = await session.execute(
        text(
            "select id, name, format, column_mapping from project_cost_import_profiles "
            "where household_id = :household_id order by name"
        ),
        {"household_id": str(household_id)},
    )
    return [CostImportMappingProfile(**row._mapping) for row in rows]


async def create_mapping_profile(
    session: AsyncSession, household_id: UUID, payload: CostImportMappingProfileInput
) -> CostImportMappingProfile:
    row = (
        await session.execute(
            text(
                "insert into project_cost_import_profiles "
                "(household_id, name, format, column_mapping) "
                "values (:household_id, :name, :format, cast(:column_mapping as jsonb)) "
                "returning id, name, format, column_mapping"
            ),
            {
                "household_id": str(household_id),
                "name": payload.name,
                "format": payload.format,
                "column_mapping": json.dumps(payload.column_mapping),
            },
        )
    ).first()
    assert row is not None
    return CostImportMappingProfile(**row._mapping)


async def delete_mapping_profile(
    session: AsyncSession, household_id: UUID, profile_id: UUID
) -> None:
    result = await session.execute(
        text(
            "delete from project_cost_import_profiles "
            "where id = :id and household_id = :household_id"
        ),
        {"id": str(profile_id), "household_id": str(household_id)},
    )
    if cast(CursorResult[Any], result).rowcount == 0:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail=COST_IMPORT_PROFILE_NOT_FOUND
        )
