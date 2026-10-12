"""Preismodul Lebensmittel: Sichtungen, Bewertungspreis, Excel-Austausch.

Die zwei Zusagen, an denen dieses Modul gemessen wird:

  1. Der bestehende Rechenkern bleibt unberuehrt. Eine vorhandene Preiszeile
     behaelt ihren Betrag, und jede Gerichtskostenschaetzung liefert nach der
     Erweiterung cent-genau dasselbe wie vorher.
  2. Angebote gehen nie in den Bewertungspreis. Ein Gericht wird nicht
     dauerhaft billiger, weil etwas diese Woche im Prospekt steht.
"""

import io
import uuid
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services import grocery_price_sheet_service as sheet_service
from app.services.grocery_observation_service import (
    Candidate,
    PriceSettings,
    resolve,
)
from app.services.grocery_price_sheet_service import (
    SHEET_NAME,
    TemplateIngredient,
    build_template_workbook,
    parse_workbook,
)
from tests.test_household_isolation import _auth_headers, _signup

client = TestClient(app)

TODAY = date(2026, 8, 5)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def household_headers() -> dict[str, str]:
    token = _signup(f"preis-{uuid.uuid4().hex}@finanzplaner.test", "correct-horse-battery-2")
    household = client.post(
        "/v1/households", json={"name": "Preismodul-Testhaushalt"}, headers=_auth_headers(token)
    ).json()
    return {**_auth_headers(token), "X-Household-Id": household["id"]}


def _candidate(
    *,
    price_minor: int,
    base_quantity: Decimal | str = "500",
    observed_on: date,
    is_promotion: bool = False,
    quality_tier: str = "standard",
    base_family: str = "masse",
    store_id: uuid.UUID | None = None,
) -> Candidate:
    quantity = Decimal(str(base_quantity))
    return Candidate(
        id=uuid.uuid4(),
        store_id=store_id,
        quality_tier=quality_tier,
        observed_on=observed_on,
        is_promotion=is_promotion,
        price_minor=price_minor,
        pack_quantity=quantity,
        pack_unit="g",
        base_family=base_family,
        unit_price_micro=int(Decimal(price_minor) * 1_000_000 / quantity),
    )


# ---------------------------------------------------------------------------
# Auswahlregel (reine Funktion)
# ---------------------------------------------------------------------------


def test_angebote_gehen_nie_in_den_bewertungspreis() -> None:
    """Die zentrale Zusage des Moduls. Das Angebot ist juenger UND guenstiger
    und gewinnt trotzdem nicht."""
    normal = _candidate(price_minor=239, observed_on=date(2026, 7, 20))
    angebot = _candidate(price_minor=179, observed_on=date(2026, 8, 4), is_promotion=True)

    entschieden = resolve([angebot, normal], settings=PriceSettings(), today=TODAY)

    assert entschieden.observation is not None
    assert entschieden.observation.price_minor == 239
    assert entschieden.reason == "zuletzt_gesehen"


def test_nur_angebotssichtungen_liefern_keinen_preis_statt_eines_falschen() -> None:
    angebot = _candidate(price_minor=179, observed_on=TODAY, is_promotion=True)

    entschieden = resolve([angebot], settings=PriceSettings(), today=TODAY)

    assert entschieden.observation is None
    assert entschieden.reason == "nur_angebote"
    assert entschieden.is_stale is True


def test_strategie_guenstigster_vergleicht_ueber_den_grundpreis() -> None:
    """1,29 fuer 500 g ist teurer als 1,99 fuer 1 kg - das sieht man nur am
    Grundpreis, und genau daran soll die Strategie entscheiden."""
    kleine_packung = _candidate(price_minor=129, base_quantity="500", observed_on=TODAY)
    grosse_packung = _candidate(price_minor=199, base_quantity="1000", observed_on=TODAY)

    entschieden = resolve(
        [kleine_packung, grosse_packung],
        settings=PriceSettings(default_strategy="cheapest"),
        today=TODAY,
    )

    assert entschieden.observation is not None
    assert entschieden.observation.price_minor == 199


def test_alte_sichtung_bleibt_gueltig_und_wird_nur_gekennzeichnet() -> None:
    """Ein alter Preis ist immer noch die beste verfuegbare Auskunft. Ihn zu
    verwerfen wuerde jedes Gericht unbewertet machen, das gestern noch eine
    Schaetzung hatte."""
    alt = _candidate(price_minor=239, observed_on=date(2025, 1, 1))

    entschieden = resolve([alt], settings=PriceSettings(max_age_days=180), today=TODAY)

    assert entschieden.observation is not None
    assert entschieden.observation.price_minor == 239
    assert entschieden.is_stale is True
    assert entschieden.reason.endswith("_alt")


def test_einheitenfamilien_werden_nicht_vermischt() -> None:
    """"500 g" gegen "1 Pkg" zu vergleichen waere geraten. Die Familie der
    juengsten Sichtung gewinnt, die andere bleibt aussen vor."""
    in_gramm = _candidate(price_minor=239, observed_on=date(2026, 8, 4), base_family="masse")
    in_packungen = _candidate(
        price_minor=99, base_quantity="1", observed_on=date(2026, 7, 1), base_family="pkg"
    )

    entschieden = resolve(
        [in_gramm, in_packungen],
        settings=PriceSettings(default_strategy="cheapest"),
        today=TODAY,
    )

    assert entschieden.observation is not None
    assert entschieden.observation.price_minor == 239


def test_bevorzugter_laden_wird_zurueckgenommen_statt_ins_leere_zu_greifen() -> None:
    """Kein Preis, nur weil man diese Woche woanders war, waere schlechter als
    ein Preis mit sichtbarer Herkunft."""
    woanders = _candidate(price_minor=239, observed_on=TODAY, store_id=uuid.uuid4())

    entschieden = resolve(
        [woanders],
        settings=PriceSettings(preferred_store_id=uuid.uuid4()),
        today=TODAY,
    )

    assert entschieden.observation is not None
    assert entschieden.observation.price_minor == 239


def test_organic_only_erfindet_keinen_bio_preis() -> None:
    konventionell = _candidate(price_minor=149, observed_on=TODAY)

    entschieden = resolve(
        [konventionell],
        settings=PriceSettings(),
        quality_preference="organic_only",
        today=TODAY,
    )

    assert entschieden.observation is None
    assert entschieden.reason == "keine_bio_sichtung"


# ---------------------------------------------------------------------------
# Excel-Vorlage
# ---------------------------------------------------------------------------


def _fill_sheet(rows: list[dict[str, object]], ingredients: list[str] | None = None) -> bytes:
    """Baut eine Vorlage und traegt Zeilen ein - wie ein Nutzer in Excel."""
    workbook_bytes = build_template_workbook(
        [TemplateIngredient(name) for name in (ingredients or [])], ["Aldi", "Rewe"]
    )
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
    sheet = workbook[SHEET_NAME]
    header = {sheet.cell(row=1, column=c).value: c for c in range(1, sheet.max_column + 1)}
    start = sheet.max_row + 1
    for offset, values in enumerate(rows):
        for label, value in values.items():
            sheet.cell(row=start + offset, column=header[label], value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_vorlage_und_einlesen_ergeben_dieselben_werte() -> None:
    data = _fill_sheet(
        [
            {
                "Zutat": "Nudeln",
                "Laden": "Aldi",
                "Packung": 500,
                "Einheit": "g",
                "Preis": 0.89,
                "gesehen am": date(2026, 8, 5),
            }
        ]
    )

    result = parse_workbook(data, today=TODAY)

    assert len(result.observations) == 1
    assert not result.problems
    gelesen = result.observations[0]
    assert gelesen.ingredient_name == "Nudeln"
    assert gelesen.ingredient_key == "nudeln"
    assert gelesen.store_name == "Aldi"
    assert gelesen.price_minor == 89
    assert gelesen.pack_quantity == Decimal("500")
    assert gelesen.base_family == "masse"
    assert gelesen.observed_on == date(2026, 8, 5)


def test_unausgefuellte_vorlagenzeilen_sind_kein_fehler() -> None:
    """Der Normalfall: 200 vorbefuellte Zutaten, zehn Preise eingetragen. Die
    restlichen 190 als Fehler zu melden wuerde die echten Fehler begraben."""
    data = _fill_sheet(
        [{"Zutat": "Nudeln", "Packung": 500, "Einheit": "g", "Preis": 0.89}],
        ingredients=["Butter", "Mehl", "Zwiebeln"],
    )

    result = parse_workbook(data, today=TODAY)

    assert len(result.observations) == 1
    assert result.untouched_rows == 3
    assert result.problems == []


def test_angebot_ohne_enddatum_wird_abgelehnt() -> None:
    """Ein Angebot ohne Ende ist kein Angebot, sondern der neue Normalpreis -
    und wuerde nach Ablauf still weiterrechnen."""
    data = _fill_sheet(
        [
            {
                "Zutat": "Butter",
                "Packung": 250,
                "Einheit": "g",
                "Preis": 1.79,
                "Angebot": "ja",
            }
        ]
    )

    result = parse_workbook(data, today=TODAY)

    assert result.observations == []
    assert len(result.problems) == 1
    assert "gueltig bis" in result.problems[0].message


def test_deutsche_zahlen_und_datumsformate_werden_gelesen() -> None:
    data = _fill_sheet(
        [
            {
                "Zutat": "Olivenöl",
                "Packung": "750",
                "Einheit": "ml",
                "Preis": "5,49 €",
                "gesehen am": "05.08.2026",
            }
        ]
    )

    result = parse_workbook(data, today=TODAY)

    assert not result.problems
    assert result.observations[0].price_minor == 549
    assert result.observations[0].observed_on == date(2026, 8, 5)


def test_identische_zeilen_in_einer_datei_werden_nur_einmal_gelesen() -> None:
    zeile = {"Zutat": "Nudeln", "Packung": 500, "Einheit": "g", "Preis": 0.89,
             "gesehen am": date(2026, 8, 5)}

    result = parse_workbook(_fill_sheet([zeile, dict(zeile)]), today=TODAY)

    assert len(result.observations) == 1
    assert len(result.problems) == 1
    assert "schon einmal identisch" in result.problems[0].message


def test_zu_grosse_datei_wird_abgewiesen(monkeypatch: pytest.MonkeyPatch) -> None:
    """Gelesen wird im Anfrage-Thread, es gibt keinen Worker dafuer. Eine
    versehentlich hochgeladene Riesendatei darf den Request nicht blockieren."""
    monkeypatch.setattr(sheet_service, "MAX_SHEET_ROWS", 5)
    zeilen = [
        {"Zutat": f"Zutat {index}", "Packung": 1, "Einheit": "kg", "Preis": 1.0}
        for index in range(20)
    ]

    with pytest.raises(sheet_service.PriceSheetError, match="Zeilen"):
        parse_workbook(_fill_sheet(zeilen), today=TODAY)


def test_spalten_duerfen_umsortiert_und_anders_beschriftet_sein() -> None:
    """Eine Datei aus fremder Hand trifft den Wortlaut der Vorlage selten
    exakt. Zugeordnet wird ueber die Beschriftung, nicht ueber die Position."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Preis (€)", "Menge", "Einheit", "Artikel", "Markt"])
    sheet.append([1.99, 1, "kg", "Mehl", "Rewe"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_workbook(buffer.getvalue(), today=TODAY)

    assert len(result.observations) == 1
    gelesen = result.observations[0]
    assert gelesen.ingredient_name == "Mehl"
    assert gelesen.price_minor == 199
    assert gelesen.store_name == "Rewe"


# ---------------------------------------------------------------------------
# Durchstich ueber die API
# ---------------------------------------------------------------------------


def test_import_fuellt_gedaechtnis_und_laesst_handeingabe_unberuehrt(
    household_headers: dict[str, str],
) -> None:
    """Die wichtigste Zusage gegenueber dem Bestand: eine von Hand getippte
    Zahl wird von keinem Import angefasst - aber die Sichtung wird trotzdem
    festgehalten, und der Nutzer erfaehrt, dass sie nicht durchschlaegt."""
    angelegt = client.post(
        "/v1/ingredient-prices",
        json={"name": "Nudeln", "price_minor": 199, "pack_quantity": "1", "pack_unit": "kg"},
        headers=household_headers,
    )
    assert angelegt.status_code == 201

    data = _fill_sheet(
        [
            {
                "Zutat": "Nudeln",
                "Laden": "Aldi",
                "Packung": 500,
                "Einheit": "g",
                "Preis": 0.89,
                "gesehen am": date(2026, 8, 5),
            }
        ]
    )
    vorschau = client.post(
        "/v1/price-sheet/preview",
        files={"file": ("Preise.xlsx", data, XLSX_MIME)},
        headers=household_headers,
    )
    assert vorschau.status_code == 200, vorschau.text
    zeilen = vorschau.json()["rows"]
    assert len(zeilen) == 1
    assert zeilen[0]["ingredient_is_new"] is False
    assert zeilen[0]["is_duplicate"] is False

    uebernahme = client.post(
        "/v1/price-observations/bulk",
        json={
            "source": "excel",
            "source_ref": "Preise.xlsx",
            "observations": [
                {
                    "ingredient_name": "Nudeln",
                    "price_minor": 89,
                    "pack_quantity": "500",
                    "pack_unit": "g",
                    "store_name": "Aldi",
                    "observed_on": "2026-08-05",
                }
            ],
        },
        headers=household_headers,
    )
    assert uebernahme.status_code == 201, uebernahme.text
    ergebnis = uebernahme.json()
    assert ergebnis["inserted"] == 1
    # Handeingabe bleibt stehen ...
    assert ergebnis["updated_valuations"] == 0
    # ... und das wird ausdruecklich gemeldet statt still hingenommen.
    assert ergebnis["blocked_by_manual"] == 1

    preise = client.get("/v1/ingredient-prices", headers=household_headers).json()
    assert [p["price_minor"] for p in preise] == [199]

    # Das Gedaechtnis ist trotzdem gefuellt.
    detail = client.get("/v1/ingredient-prices/nudeln/detail", headers=household_headers)
    assert detail.status_code == 200
    assert len(detail.json()["history"]) == 1


def test_umstellung_auf_ableitung_uebernimmt_die_sichtung(
    household_headers: dict[str, str],
) -> None:
    client.post(
        "/v1/ingredient-prices",
        json={"name": "Butter", "price_minor": 999, "pack_quantity": "250", "pack_unit": "g"},
        headers=household_headers,
    )
    client.post(
        "/v1/price-observations",
        json={
            "ingredient_name": "Butter",
            "price_minor": 229,
            "pack_quantity": "250",
            "pack_unit": "g",
            "store_name": "Aldi",
            "observed_on": "2026-08-05",
        },
        headers=household_headers,
    )

    umgestellt = client.patch(
        "/v1/ingredient-prices/butter/valuation",
        json={"price_source": "derived", "price_strategy": "latest"},
        headers=household_headers,
    )
    assert umgestellt.status_code == 200, umgestellt.text
    assert umgestellt.json()["price_source"] == "derived"

    preise = client.get("/v1/ingredient-prices", headers=household_headers).json()
    assert [p["price_minor"] for p in preise] == [229]


def test_doppelter_import_derselben_datei_ist_folgenlos(
    household_headers: dict[str, str],
) -> None:
    posten = {
        "ingredient_name": "Mehl",
        "price_minor": 99,
        "pack_quantity": "1",
        "pack_unit": "kg",
        "store_name": "Aldi",
        "observed_on": "2026-08-05",
    }
    erster = client.post(
        "/v1/price-observations/bulk",
        json={"observations": [posten]},
        headers=household_headers,
    ).json()
    zweiter = client.post(
        "/v1/price-observations/bulk",
        json={"observations": [posten]},
        headers=household_headers,
    ).json()

    assert erster["inserted"] == 1
    assert erster["created_price_rows"] == 1
    assert zweiter["inserted"] == 0
    assert zweiter["skipped_duplicates"] == 1

    detail = client.get("/v1/ingredient-prices/mehl/detail", headers=household_headers).json()
    assert len(detail["history"]) == 1


def test_neue_zutat_wird_angelegt_und_sofort_bewertet(
    household_headers: dict[str, str],
) -> None:
    """Eine Sichtung nuetzt nichts, solange der Rechenkern keine Preiszeile
    dieses Namens findet."""
    ergebnis = client.post(
        "/v1/price-observations/bulk",
        json={
            "observations": [
                {
                    "ingredient_name": "Haferflocken",
                    "price_minor": 79,
                    "pack_quantity": "500",
                    "pack_unit": "g",
                    "store_name": "Aldi",
                    "observed_on": "2026-08-05",
                }
            ]
        },
        headers=household_headers,
    ).json()

    assert ergebnis["created_price_rows"] == 1
    assert ergebnis["updated_valuations"] == 1

    preise = client.get("/v1/ingredient-prices", headers=household_headers).json()
    haferflocken = next(p for p in preise if p["name"] == "Haferflocken")
    assert haferflocken["price_minor"] == 79


def test_preismatrix_zeigt_laeden_nebeneinander(household_headers: dict[str, str]) -> None:
    """Der Fall, um den es dem Nutzer geht: das Angebot bei Rewe ist teurer
    als der Normalpreis bei Aldi - sichtbar nur ueber den Grundpreis."""
    for laden, preis, angebot in (("Rewe", 699, True), ("Aldi", 549, False)):
        posten: dict[str, object] = {
            "ingredient_name": "Gerolsteiner Medium",
            "price_minor": preis,
            "pack_quantity": "12",
            "pack_unit": "Stk",
            "store_name": laden,
            "observed_on": "2026-08-05",
            "is_promotion": angebot,
        }
        if angebot:
            posten["offer_valid_to"] = "2026-08-08"
        antwort = client.post(
            "/v1/price-observations", json=posten, headers=household_headers
        )
        assert antwort.status_code == 201, antwort.text

    detail = client.get(
        "/v1/ingredient-prices/gerolsteiner medium/detail", headers=household_headers
    ).json()

    zellen = {zelle["store_label"]: zelle for zelle in detail["matrix"]}
    assert zellen["Rewe"]["unit_price_micro"] > zellen["Aldi"]["unit_price_micro"]
    assert zellen["Rewe"]["is_promotion"] is True
    # Der Bewertungspreis nimmt trotzdem den Normalpreis.
    preise = client.get("/v1/ingredient-prices", headers=household_headers).json()
    gerolsteiner = next(p for p in preise if p["name"] == "Gerolsteiner Medium")
    assert gerolsteiner["price_minor"] == 549


def test_geloeschte_sichtung_faellt_auf_die_naechste_zurueck(
    household_headers: dict[str, str],
) -> None:
    for preis, tag in ((199, "2026-07-01"), (179, "2026-08-05")):
        client.post(
            "/v1/price-observations",
            json={
                "ingredient_name": "Reis",
                "price_minor": preis,
                "pack_quantity": "1",
                "pack_unit": "kg",
                "observed_on": tag,
            },
            headers=household_headers,
        )

    detail = client.get("/v1/ingredient-prices/reis/detail", headers=household_headers).json()
    assert detail["valuation"]["price_source"] == "derived"
    juengste = detail["history"][0]
    assert juengste["price_minor"] == 179

    geloescht = client.delete(
        f"/v1/price-observations/{juengste['id']}", headers=household_headers
    )
    assert geloescht.status_code == 204

    preise = client.get("/v1/ingredient-prices", headers=household_headers).json()
    assert [p["price_minor"] for p in preise] == [199]


def test_vorlage_enthaelt_bekannte_zutaten_und_laesst_sich_wieder_einlesen(
    household_headers: dict[str, str],
) -> None:
    client.post(
        "/v1/ingredient-prices",
        json={"name": "Zwiebeln", "price_minor": 149, "pack_quantity": "1", "pack_unit": "kg"},
        headers=household_headers,
    )

    antwort = client.get("/v1/price-sheet/template", headers=household_headers)

    assert antwort.status_code == 200
    assert "attachment" in antwort.headers["content-disposition"]
    result = parse_workbook(antwort.content, today=TODAY)
    # Die Zutat steht drin, aber ohne Preis - eine Vorlage bringt keine Zahlen
    # mit, die niemand gesehen hat.
    assert result.observations == []
    assert result.untouched_rows == 1


def test_gerichtskosten_bleiben_nach_der_erweiterung_unveraendert(
    household_headers: dict[str, str],
) -> None:
    """Abnahmekriterium der Migration: am Tag 1 springt keine Zahl."""
    client.post(
        "/v1/ingredient-prices",
        json={"name": "Nudeln", "price_minor": 199, "pack_quantity": "1", "pack_unit": "kg"},
        headers=household_headers,
    )
    gericht = client.post(
        "/v1/dishes", json={"name": "Nudelauflauf", "default_servings": 4},
        headers=household_headers,
    ).json()
    client.post(
        "/v1/dish-ingredients",
        json={"dish_id": gericht["id"], "name": "Nudeln", "quantity": "500", "unit": "g"},
        headers=household_headers,
    )

    gerichte = client.get("/v1/dishes", headers=household_headers).json()
    auflauf = next(d for d in gerichte if d["name"] == "Nudelauflauf")

    # 1,99 EUR je kg, 500 g im Rezept -> 1,00 EUR (kaufmaennisch gerundet).
    assert auflauf["computed_cost_minor"] == 100
    assert auflauf["unpriced_ingredients"] == 0
