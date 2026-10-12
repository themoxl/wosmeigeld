"""XLSX-/CSV-Export eines Vorhaben-Szenarios (Plan Abschnitt 10.3, P4)."""

import io
import uuid
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from tests.test_household_isolation import _auth_headers, _signup

client = TestClient(app)

_EXPECTED_SHEETS = [
    "Zusammenfassung", "Kosten", "Finanzierung", "Tilgungsplaene",
    "Cashflow", "Foerderungen", "Annahmen", "Ist-Zuordnungen",
]


@pytest.fixture()
def household_headers() -> dict[str, str]:
    token = _signup(f"export-{uuid.uuid4().hex}@finanzplaner.test", "correct-horse-battery-8")
    household = client.post(
        "/v1/households", json={"name": "Export-Testhaushalt"}, headers=_auth_headers(token)
    ).json()
    return {**_auth_headers(token), "X-Household-Id": household["id"]}


def _create_project(headers: dict[str, str]) -> dict[str, Any]:
    response = client.post(
        "/v1/projects",
        json={"name": "Export-Projekt", "project_type": "home_build"},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    return dict(response.json())


def _baseline_scenario(headers: dict[str, str], project_id: str) -> dict[str, Any]:
    scenarios = client.get(f"/v1/projects/{project_id}/scenarios", headers=headers).json()
    return dict(scenarios[0])


def test_export_xlsx_has_all_eight_sheets_in_order(household_headers: dict[str, str]) -> None:
    project = _create_project(household_headers)
    scenario = _baseline_scenario(household_headers, project["id"])
    section = client.post(
        f"/v1/projects/{project['id']}/scenarios/{scenario['id']}/cost-sections",
        json={"name": "Rohbau", "cost_class": "direct"},
        headers=household_headers,
    ).json()
    client.post(
        f"/v1/projects/{project['id']}/scenarios/{scenario['id']}/cost-items",
        json={
            "name": "Fundament", "section_id": section["id"],
            "calculation_mode": "fixed", "estimate_manual_minor": 12_000_00,
        },
        headers=household_headers,
    )

    response = client.get(
        f"/v1/projects/{project['id']}/scenarios/{scenario['id']}/export.xlsx",
        headers=household_headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == _EXPECTED_SHEETS

    costs_sheet = workbook["Kosten"]
    rows = list(costs_sheet.iter_rows(values_only=True))
    header_row_index = next(i for i, r in enumerate(rows) if r and r[0] == "Position")
    data_row = rows[header_row_index + 1]
    assert data_row[0] == "Fundament"
    assert data_row[1] == "Rohbau"
    estimate_col = rows[header_row_index].index("Schaetzung")
    # openpyxl liest eine ganzzahlige Zahl als int zurueck, nicht als float -
    # numerisch ausschlaggebend ist nur, dass es KEIN String ist (Plan 10.3:
    # "Zahlen ... nicht als lokalisierte Textbetraege").
    assert data_row[estimate_col] == 12_000
    assert isinstance(data_row[estimate_col], int | float)


def test_export_csv_matches_cost_items(household_headers: dict[str, str]) -> None:
    project = _create_project(household_headers)
    scenario = _baseline_scenario(household_headers, project["id"])
    section = client.post(
        f"/v1/projects/{project['id']}/scenarios/{scenario['id']}/cost-sections",
        json={"name": "Rohbau", "cost_class": "direct"},
        headers=household_headers,
    ).json()
    client.post(
        f"/v1/projects/{project['id']}/scenarios/{scenario['id']}/cost-items",
        json={
            "name": "Fundament", "section_id": section["id"],
            "calculation_mode": "fixed", "estimate_manual_minor": 12_000_00,
        },
        headers=household_headers,
    )

    response = client.get(
        f"/v1/projects/{project['id']}/scenarios/{scenario['id']}/export/cost-items.csv",
        headers=household_headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    text = response.content.decode("utf-8")
    assert text.startswith("﻿")
    assert "Fundament" in text
    assert "12000.0" in text or "12000.00" in text


def test_export_requires_scenario_belongs_to_project(household_headers: dict[str, str]) -> None:
    project = _create_project(household_headers)
    other_project = _create_project(household_headers)
    scenario = _baseline_scenario(household_headers, project["id"])

    response = client.get(
        f"/v1/projects/{other_project['id']}/scenarios/{scenario['id']}/export/cost-items.csv",
        headers=household_headers,
    )
    assert response.status_code == 404
