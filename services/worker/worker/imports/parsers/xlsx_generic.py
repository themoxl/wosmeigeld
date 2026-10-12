"""Generischer XLSX-Kontoauszug-Parser."""

import io
from collections.abc import Iterator, Mapping
from datetime import date, datetime
from itertools import chain

import openpyxl

from worker.imports.canonical import (
    CanonicalTransaction,
    ImportParseError,
    parse_amount_to_minor_units,
)
from worker.imports.parsers.base import DetectionResult
from worker.imports.parsers.csv_generic import (
    AMOUNT_ALIASES,
    DATE_ALIASES,
    DATE_FORMATS,
    PAYEE_ALIASES,
    PURPOSE_ALIASES,
    REFERENCE_ALIASES,
    STATUS_ALIASES,
    _find_column,
    is_pending_status,
)

# Wie weit oben in der Datei nach einer Kopfzeile gesucht wird - dieselbe
# Spanne, die `detect` benutzt.
HEADER_SCAN_ROWS = 20


def _trim_trailing_empty(headers: list[str]) -> list[str]:
    """Leere Spalten am Zeilenende abschneiden.

    Tabellenprogramme melden gern mehr Spalten, als das Blatt Ueberschriften
    hat. Ohne diesen Schnitt stuenden in detected_columns Phantomspalten, die
    im Mappingdialog zur Auswahl angeboten wuerden."""
    trimmed = list(headers)
    while trimmed and not trimmed[-1]:
        trimmed.pop()
    return trimmed


def _unique_headers(headers: list[str]) -> list[str]:
    """Doppelte Spaltenueberschriften unterscheidbar machen.

    Der easybank-Export fuehrt 'Buchungsdatum' zweimal (Umsatz- und
    Buchungstag). Als Schluessel eines dict verschluckt die zweite Spalte die
    erste - in raw_data stand danach nur noch einer der beiden Werte."""
    seen: dict[str, int] = {}
    result: list[str] = []
    for header in headers:
        count = seen.get(header, 0)
        seen[header] = count + 1
        result.append(header if count == 0 else f"{header} ({count + 1})")
    return result


def _parse_cell_date(val: object) -> date:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        cleaned = val.strip()
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(cleaned, fmt).date()
            except ValueError:
                continue
    raise ImportParseError(f"unbekanntes Datumsformat in XLSX: {val!r}")


def _extract_rows(
    rows_iter: Iterator[tuple[object, ...]],
    headers: list[str],
    date_idx: int,
    amount_idx: int,
    payee_idx: int | None,
    purpose_idx: int | None,
    ref_idx: int | None,
    status_idx: int | None = None,
) -> list[CanonicalTransaction]:
    # Fuehrt der Export nur eine Beschreibungsspalte, traegt sie Empfaenger und
    # Zweck zugleich - dann bleibt der Zweck leer statt den Empfaenger wortgleich
    # zu wiederholen.
    if purpose_idx is not None and purpose_idx == payee_idx:
        purpose_idx = None

    unique_headers = _unique_headers(headers)
    results: list[CanonicalTransaction] = []
    for row in rows_iter:
        if not row or all(v is None for v in row):
            continue
        if status_idx is not None and status_idx < len(row) and is_pending_status(row[status_idx]):
            continue
        date_val = row[date_idx] if date_idx < len(row) else None
        amount_val = row[amount_idx] if amount_idx < len(row) else None
        if date_val is None or amount_val is None:
            continue

        try:
            booked_date = _parse_cell_date(date_val)
            amount_minor = parse_amount_to_minor_units(str(amount_val))
        except ImportParseError:
            continue

        payee = (
            str(row[payee_idx]).strip()
            if payee_idx is not None and payee_idx < len(row) and row[payee_idx]
            else None
        )
        purpose = (
            str(row[purpose_idx]).strip()
            if purpose_idx is not None and purpose_idx < len(row) and row[purpose_idx]
            else None
        )
        # Der Index-Vergleich oben faengt nur die eine Spalte, die beide Rollen
        # traegt. Fuehrt der Export ZWEI Spalten mit wortgleichem Inhalt (der
        # easybank-Export hat "Beschreibung" und "Details" mit identischem
        # Text), stehen Empfaenger und Zweck sonst doppelt in der Buchung.
        if purpose is not None and purpose == payee:
            purpose = None
        ref = (
            str(row[ref_idx]).strip()
            if ref_idx is not None and ref_idx < len(row) and row[ref_idx]
            else None
        )

        results.append(
            CanonicalTransaction(
                booked_date=booked_date,
                amount_minor=amount_minor,
                currency="EUR",
                payee_raw=payee,
                purpose=purpose,
                external_reference=ref,
                raw={
                    unique_headers[i]: str(row[i])
                    for i in range(min(len(unique_headers), len(row)))
                },
            )
        )

    return results


class GenericXlsxParser:
    format_name = "xlsx_generic"

    def detect(self, file_bytes: bytes, filename: str) -> DetectionResult:
        if not filename.lower().endswith(".xlsx"):
            return DetectionResult(False, self.format_name, 0.0, "keine .xlsx-Endung")
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
            sheet = wb.active
            if sheet is None:
                return DetectionResult(False, self.format_name, 0.0, "kein aktives Blatt")
            for row in sheet.iter_rows(max_row=20, values_only=True):
                headers = [str(c).strip() for c in row if c is not None]
                has_date = _find_column(headers, DATE_ALIASES) is not None
                has_amount = _find_column(headers, AMOUNT_ALIASES) is not None
                if has_date and has_amount:
                    wb.close()
                    return DetectionResult(True, self.format_name, 0.7, "XLSX Datum/Betrag erkannt")
            wb.close()
            return DetectionResult(
                True, self.format_name, 0.3, "XLSX erkannt, Spalten-Mapping benoetigt"
            )
        except Exception as exc:
            return DetectionResult(False, self.format_name, 0.0, f"XLSX-Ladefehler: {exc}")

    def parse(self, file_bytes: bytes) -> list[CanonicalTransaction]:
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
        sheet = wb.active
        if sheet is None:
            wb.close()
            return []

        rows_iter = sheet.iter_rows(values_only=True)
        headers: list[str] = []
        date_col: str | None = None
        amount_col: str | None = None

        # Preamble-Zeilen durchsuchen, bis Spaltenkopf mit Datum/Betrag gefunden wird
        for row in rows_iter:
            if not row:
                continue
            candidate_headers = [str(h).strip() if h is not None else "" for h in row]
            date_candidate = _find_column(candidate_headers, DATE_ALIASES)
            amount_candidate = _find_column(candidate_headers, AMOUNT_ALIASES)
            if date_candidate and amount_candidate:
                headers = candidate_headers
                date_col = date_candidate
                amount_col = amount_candidate
                break

        if not date_col or not amount_col:
            wb.close()
            raise ImportParseError("Datums- oder Betragsspalte in XLSX nicht erkannt.")

        date_idx = headers.index(date_col)
        amount_idx = headers.index(amount_col)

        payee_col = _find_column(headers, PAYEE_ALIASES)
        payee_idx = headers.index(payee_col) if payee_col else None

        purpose_col = _find_column(headers, PURPOSE_ALIASES)
        purpose_idx = headers.index(purpose_col) if purpose_col else None

        ref_col = _find_column(headers, REFERENCE_ALIASES)
        ref_idx = headers.index(ref_col) if ref_col else None

        status_col = _find_column(headers, STATUS_ALIASES)
        status_idx = headers.index(status_col) if status_col else None

        results = _extract_rows(
            rows_iter,
            headers,
            date_idx,
            amount_idx,
            payee_idx,
            purpose_idx,
            ref_idx,
            status_idx,
        )
        wb.close()
        return results

    def read_raw_rows(
        self, file_bytes: bytes, limit: int | None = None
    ) -> tuple[list[str], list[dict[str, str]]]:
        """Kopfzeile und Rohzeilen, ohne irgendeine Spalte zu deuten (siehe
        GenericCsvParser.read_raw_rows).

        Die Kopfzeile wird zuerst wie in `parse` gesucht (erste Zeile mit
        Datums- UND Betragsspalte), damit Praeambelzeilen nicht als
        Spaltennamen durchgehen. Erst wenn es die nicht gibt - genau der Fall,
        fuer den es 'needs_mapping' ueberhaupt gibt -, gilt die erste nicht
        leere Zeile als Kopfzeile."""
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
        sheet = wb.active
        if sheet is None:
            wb.close()
            return [], []

        rows_iter = sheet.iter_rows(values_only=True)
        headers: list[str] = []
        # Vorlauf, bis die Kopfzeile feststeht: der Iterator eines im
        # read_only-Modus geoeffneten Blatts laesst sich nicht zuruecksetzen.
        preamble: list[tuple[object, ...]] = []
        for row in rows_iter:
            if not row or all(v is None for v in row):
                continue
            candidate = [str(h).strip() if h is not None else "" for h in row]
            if _find_column(candidate, DATE_ALIASES) and _find_column(candidate, AMOUNT_ALIASES):
                headers = candidate
                preamble = []
                break
            preamble.append(row)
            if len(preamble) >= HEADER_SCAN_ROWS:
                break

        data_iter: Iterator[tuple[object, ...]] = rows_iter
        if not headers:
            if not preamble:
                wb.close()
                return [], []
            headers = [str(h).strip() if h is not None else "" for h in preamble[0]]
            data_iter = chain(preamble[1:], rows_iter)

        headers = _unique_headers(_trim_trailing_empty(headers))
        rows: list[dict[str, str]] = []
        for row in data_iter:
            if limit is not None and len(rows) >= limit:
                break
            if not row or all(v is None for v in row):
                continue
            rows.append(
                {
                    headers[i]: "" if row[i] is None else str(row[i])
                    for i in range(min(len(headers), len(row)))
                }
            )
        wb.close()
        return headers, rows

    def parse_with_mapping(
        self, file_bytes: bytes, mapping: Mapping[str, str | None]
    ) -> list[CanonicalTransaction]:
        date_col = mapping.get("booked_date")
        amount_col = mapping.get("amount")
        if not date_col or not amount_col:
            raise ImportParseError("Mapping benoetigt mindestens Datum und Betrag.")

        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
        sheet = wb.active
        if sheet is None:
            wb.close()
            return []

        rows_iter = sheet.iter_rows(values_only=True)
        headers: list[str] = []

        for row in rows_iter:
            if not row:
                continue
            candidate_headers = [str(h).strip() if h is not None else "" for h in row]
            if date_col in candidate_headers and amount_col in candidate_headers:
                headers = candidate_headers
                break

        if not headers:
            wb.close()
            raise ImportParseError(
                f"Mapping-Spalten {date_col!r}/{amount_col!r} nicht in XLSX gefunden."
            )

        date_idx = headers.index(date_col)
        amount_idx = headers.index(amount_col)

        payee_col = mapping.get("payee")
        payee_idx = headers.index(payee_col) if payee_col and payee_col in headers else None

        purpose_col = mapping.get("purpose")
        purpose_idx = (
            headers.index(purpose_col) if purpose_col and purpose_col in headers else None
        )

        ref_col = mapping.get("reference")
        ref_idx = headers.index(ref_col) if ref_col and ref_col in headers else None

        # Ein Mapping-Profil muss die Statusspalte nicht kennen; ist sie im
        # Blatt vorhanden, werden Vormerkungen trotzdem uebersprungen.
        status_col = mapping.get("status") or _find_column(headers, STATUS_ALIASES)
        status_idx = headers.index(status_col) if status_col and status_col in headers else None

        results = _extract_rows(
            rows_iter,
            headers,
            date_idx,
            amount_idx,
            payee_idx,
            purpose_idx,
            ref_idx,
            status_idx,
        )
        wb.close()
        return results
